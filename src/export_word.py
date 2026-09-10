"""Exporta la memoria a Word según la plantilla y las normas de presentación.

Reutiliza Formato/plantilla-portada-tfm.docx (logos y campos). El cuerpo sigue
Calibri 11, interlineado 1,5, 6 pt entre párrafos, texto justificado, pies de
figura a 8 pt centrados, capítulos en página nueva y numeración inferior derecha
a partir de la página 4 (portada, blanco y agradecimientos cuentan y no numeran).
"""

from __future__ import annotations

import re
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from assemble_report import (
    CHAPTERS,
    REPORTS_DIR,
    SUBTITLE,
    TITLE,
    embed_figures,
)

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Formato" / "plantilla-portada-tfm.docx"
OUTPUT_PATH = REPORTS_DIR / "TFM_memoria.docx"
FIGURES_DIR = REPORTS_DIR / "figures"

# Completar estos tres campos de la portada oficial antes de la entrega.
COVER = {
    "master": "MÁSTER EN [indicar titulación y modalidad]",
    "author": "[Nombre y apellidos]",
    "tutor": "[Nombre y apellidos del tutor/a]",
    "place_date": "- septiembre de 2026 -",
}

HEADING_RE = re.compile(r"^(#{1,3}) (.+)$")
IMAGE_RE = re.compile(r"^!\[.*?\]\((.+?)\)$")
CAPTION_RE = re.compile(r"^\*{0,3}Figura ([^*\s]+)\.{0,1}\*{0,3}\s*\*{0,1}(.*)\*{0,1}$")
LIST_RE = re.compile(r"^(\s*)([-*]|\d+\.) (.+)$")
REF_ENTRY_RE = re.compile(r"^[A-ZÁÉÍÓÚÜÑ].*\(\d{4}[a-z]?\).*")

INLINE_TOKEN = re.compile(
    r"(`[^`]+`)"
    r"|(\*\*\*[^*]+\*\*\*)"
    r"|(\*\*[^*]+\*\*)"
    r"|(\*[^*\s][^*]*\*)"
    r"|(\$[^$`]{1,80}\$)"
    r"|(\[[^\]]+\]\([^)]+\))"
)

SUB_TRANS = str.maketrans(
    "0123456789+-=()aeioruvx",
    "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑᵢₒᵣᵤᵥₓ",
)
SUP_TRANS = str.maketrans(
    "0123456789+-=()ni",
    "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱ",
)

LATEX_SIMPLE = [
    (r"\cdot", "·"),
    (r"\times", "×"),
    (r"\geq", "≥"),
    (r"\leq", "≤"),
    (r"\neq", "≠"),
    (r"\approx", "≈"),
    (r"\infty", "∞"),
    (r"\pm", "±"),
    (r"\ldots", "…"),
    (r"\dots", "…"),
    (r"\circ", "°"),
    (r"\in", "∈"),
    (r"\subset", "⊂"),
    (r"\forall", "∀"),
    (r"\partial", "∂"),
    (r"\ell", "ℓ"),
    (r"\alpha", "α"),
    (r"\beta", "β"),
    (r"\gamma", "γ"),
    (r"\delta", "δ"),
    (r"\eta", "η"),
    (r"\theta", "θ"),
    (r"\lambda", "λ"),
    (r"\mu", "μ"),
    (r"\pi", "π"),
    (r"\sigma", "σ"),
    (r"\phi", "ϕ"),
    (r"\varphi", "φ"),
    (r"\omega", "ω"),
    (r"\Phi", "Φ"),
    (r"\mathbb{R}", "ℝ"),
    (r"\mathbb{1}", "𝟙"),
    (r"\mathbf{x}", "x"),
    (r"\max", "max"),
    (r"\min", "min"),
    (r"\sum", "∑"),
    (r"\prod", "∏"),
    (r"\hat", ""),
    (r"\bar", ""),
    (r"\tilde", ""),
    (r"\mathbf", ""),
    (r"\mathrm", ""),
    (r"\mathbb", ""),
    (r"\mathcal", ""),
    (r"\operatorname", ""),
    (r"\left", ""),
    (r"\right", ""),
    (r"\;", " "),
    (r"\,", " "),
    (r"\!", ""),
    (r"\ ", " "),
]


def latex_to_text(expr: str) -> str:
    text = expr.strip().strip("$").replace("{,}", ",").replace("`", "")
    text = re.sub(
        r"\\frac\{([^{}]+)\}\{([^{}]+)\}",
        lambda m: f"({m.group(1)})/({m.group(2)})",
        text,
    )
    text = re.sub(
        r"_\{([^{}]+)\}",
        lambda m: m.group(1).translate(SUB_TRANS),
        text,
    )
    text = re.sub(
        r"\^\{([^{}]+)\}",
        lambda m: m.group(1).translate(SUP_TRANS),
        text,
    )
    text = re.sub(r"_([A-Za-z0-9])", lambda m: m.group(1).translate(SUB_TRANS), text)
    text = re.sub(r"\^([A-Za-z0-9])", lambda m: m.group(1).translate(SUP_TRANS), text)
    for src, dst in LATEX_SIMPLE:
        text = text.replace(src, dst)
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\\[A-Za-z]+", "", text)
    return re.sub(r" +", " ", text).strip()


def sanitize_md(text: str) -> str:
    """Elimina restos de Markdown/LaTeX que no deben verse en papel."""
    text = text.replace("{,}", ",")
    text = re.sub(r"(\$[A-Za-z][^=$`]*=\s*0,\d+)`", r"\1$", text)
    text = re.sub(r"\*\*`([^`]+)`\*\*", r"**\1**", text)
    text = re.sub(r"`\*\*([^*]+)\*\*`", r"**\1**", text)
    text = re.sub(r"`(?:reports|src|data)/[^`]+`", "", text)
    return text


def set_run_font(run, name: str, size_pt: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def apply_style_font(style, name: str, size_pt: float, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def set_paragraph_body(paragraph: Paragraph, *, justify: bool = True) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(6)
    fmt.space_before = Pt(0)
    fmt.widow_control = True
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def add_page_number(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, "Calibri", 11)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run = paragraph.add_run()
    set_run_font(run, "Calibri", 11)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run._r.append(instr)

    run = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)

    run = paragraph.add_run(" ")
    set_run_font(run, "Calibri", 11)

    run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)


def set_section_start(section, start: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def enable_update_fields(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def collect_headings() -> list[tuple[int, str]]:
    items: list[tuple[int, str]] = []
    for name in CHAPTERS:
        text = (REPORTS_DIR / name).read_text(encoding="utf-8")
        for line in text.splitlines():
            heading = HEADING_RE.match(line.strip())
            if heading:
                items.append((len(heading.group(1)), strip_md_inline(heading.group(2).strip())))
    return items


def add_toc(document: Document) -> None:
    """Inserta un TOC de Word (niveles 1–3) con números de página y líderes de tabulación."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    run._r.append(instr)

    run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)

    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def ensure_heading_style(document: Document, name: str, outline_level: int):
    try:
        style = document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
    p_pr = style.element.get_or_add_pPr()
    for child in p_pr.findall(qn("w:outlineLvl")):
        p_pr.remove(child)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(outline_level))
    p_pr.append(outline)
    return style


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    apply_style_font(normal, "Calibri", 11, False)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    fmt = normal.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(6)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_specs = {
        1: (18, True, Pt(0), Pt(12)),
        2: (16, True, Pt(12), Pt(6)),
        3: (14, True, Pt(10), Pt(6)),
    }
    for level, (size, bold, before, after) in heading_specs.items():
        style = ensure_heading_style(document, f"Heading {level}", level - 1)
        apply_style_font(style, "Calibri", size, bold)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = before
        style.paragraph_format.space_after = after
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if level == 1:
            style.paragraph_format.page_break_before = True


def add_formatted_runs(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    text = sanitize_md(text)
    pos = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, "Calibri", 11, bold)
            run.italic = italic
        token = match.group(0)
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Calibri", 11, bold)
            run.italic = True
        elif token.startswith("***") and token.endswith("***"):
            inner = sanitize_md(token[3:-3]).replace("`", "")
            run = paragraph.add_run(inner)
            set_run_font(run, "Calibri", 11, True)
            run.italic = True
        elif token.startswith("**") and token.endswith("**"):
            inner = token[2:-2]
            if "`" in inner or "$" in inner:
                add_formatted_runs(paragraph, inner, bold=True, italic=italic)
            else:
                run = paragraph.add_run(inner)
                set_run_font(run, "Calibri", 11, True)
                run.italic = italic
        elif token.startswith("*") and token.endswith("*"):
            inner = token[1:-1]
            if "`" in inner or "$" in inner:
                add_formatted_runs(paragraph, inner, bold=bold, italic=True)
            else:
                run = paragraph.add_run(inner)
                set_run_font(run, "Calibri", 11, bold)
                run.italic = True
        elif token.startswith("$") and token.endswith("$"):
            run = paragraph.add_run(latex_to_text(token))
            set_run_font(run, "Calibri", 11, bold)
            run.italic = True
        elif token.startswith("[") and "](" in token:
            label = token[1 : token.index("]")]
            run = paragraph.add_run(label)
            set_run_font(run, "Calibri", 11, bold)
            run.italic = italic
        else:
            run = paragraph.add_run(token.replace("`", "").replace("$", ""))
            set_run_font(run, "Calibri", 11, bold)
            run.italic = italic
        pos = match.end()
    rest = text[pos:].replace("$", "").replace("`", "")
    if rest:
        run = paragraph.add_run(rest)
        set_run_font(run, "Calibri", 11, bold)
        run.italic = italic


def strip_md_inline(text: str) -> str:
    text = sanitize_md(text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*\*([^*]+)\*\*\*", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\$([^$]+)\$", lambda m: latex_to_text(m.group(1)), text)
    return text.replace("$", "").replace("`", "")


def add_body_paragraph(
    document: Document,
    text: str,
    *,
    hanging: bool = False,
    first_line: bool = False,
) -> Paragraph:
    paragraph = document.add_paragraph()
    set_paragraph_body(paragraph)
    if hanging:
        paragraph.paragraph_format.left_indent = Cm(1.25)
        paragraph.paragraph_format.first_line_indent = Cm(-1.25)
    elif first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0)
    add_formatted_runs(paragraph, text)
    return paragraph


def add_heading(document: Document, text: str, level: int) -> None:
    heading = document.add_heading(strip_md_inline(text), level=level)
    set_paragraph_body(heading, justify=False)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        set_run_font(run, "Calibri", {1: 18, 2: 16, 3: 14}[level], True)
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(3)
    fmt.space_after = Pt(12)
    fmt.line_spacing = 1.0
    run = paragraph.add_run(sanitize_md(text).replace("`", ""))
    set_run_font(run, "Calibri", 8)
    run.italic = True


def add_image(document: Document, rel_path: str) -> None:
    path = (REPORTS_DIR / rel_path).resolve()
    if not path.exists():
        path = (FIGURES_DIR / Path(rel_path).name).resolve()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    if path.exists():
        run.add_picture(str(path), width=Cm(15.0))
    else:
        set_run_font(run, "Calibri", 11)
        run.italic = True
        run.text = f"[Figura no encontrada: {rel_path}]"


def add_equation(document: Document, expr: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(6)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = paragraph.add_run(latex_to_text(expr))
    set_run_font(run, "Calibri", 11)
    run.italic = True


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells if cell != "")


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(document: Document, lines: list[str]) -> None:
    rows = [parse_table_row(line) for line in lines if line.strip()]
    rows = [row for row in rows if not is_separator_row(row)]
    if not rows:
        return
    ncols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0
            value = row[j] if j < len(row) else ""
            if i == 0:
                run = paragraph.add_run(strip_md_inline(value))
                set_run_font(run, "Calibri", 10, True)
            else:
                add_formatted_runs(paragraph, value)
                for run in paragraph.runs:
                    if run.font.size is None or run.font.size.pt != 10:
                        if run.font.name != "Consolas":
                            set_run_font(run, "Calibri", 10)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def add_list_item(document: Document, text: str, ordered: bool, level: int) -> None:
    style_name = "List Number" if ordered else "List Bullet"
    try:
        paragraph = document.add_paragraph(style=style_name)
    except KeyError:
        paragraph = document.add_paragraph()
        prefix = "• " if not ordered else ""
        text = prefix + text
    set_paragraph_body(paragraph, justify=True)
    paragraph.paragraph_format.left_indent = Cm(1.0 + 0.75 * level)
    paragraph.paragraph_format.first_line_indent = Cm(-0.5)
    add_formatted_runs(paragraph, text)


def iter_blocks(text: str):
    lines = text.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue
        if stripped in {"---", "***"}:
            i += 1
            continue
        if stripped.startswith("*Fuente:") and re.search(r"reports/|src/|data/", stripped):
            i += 1
            continue
        if stripped.startswith("*Tabla completa"):
            i += 1
            continue
        if re.fullmatch(r"`(?:reports|src|data)/[^`]+`", stripped):
            i += 1
            continue
        if stripped.startswith("<!--"):
            i += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            yield "heading", (len(heading.group(1)), heading.group(2).strip())
            i += 1
            continue

        image = IMAGE_RE.match(stripped)
        if image:
            yield "image", image.group(1)
            i += 1
            if i < n and not lines[i].strip():
                i += 1
            if i < n:
                cap = lines[i].strip()
                if cap.startswith("***Figura") or cap.startswith("*Figura") or cap.startswith("Figura"):
                    yield "caption", cap
                    i += 1
            continue

        if stripped.startswith("$$"):
            expr = stripped[2:].strip()
            i += 1
            parts = [expr] if expr else []
            while i < n and lines[i].strip() != "$$":
                parts.append(lines[i].strip())
                i += 1
            if i < n and lines[i].strip() == "$$":
                i += 1
            yield "equation", " ".join(p for p in parts if p)
            continue

        if stripped.startswith("|"):
            table = [stripped]
            i += 1
            while i < n and lines[i].strip().startswith("|"):
                table.append(lines[i].strip())
                i += 1
            yield "table", table
            continue

        list_match = LIST_RE.match(line)
        if list_match:
            indent, marker, content = list_match.groups()
            level = len(indent.replace("\t", "    ")) // 2
            pieces = [content.strip()]
            i += 1
            while i < n:
                nxt = lines[i]
                if not nxt.strip():
                    break
                if LIST_RE.match(nxt) or HEADING_RE.match(nxt.strip()) or nxt.strip().startswith("|"):
                    break
                if nxt.startswith("  ") or nxt.startswith("\t"):
                    pieces.append(nxt.strip())
                    i += 1
                    continue
                break
            yield "list", (marker.endswith("."), level, " ".join(pieces))
            continue

        if stripped.startswith(">"):
            pieces = [stripped.lstrip("> ").strip()]
            i += 1
            while i < n and lines[i].strip().startswith(">"):
                pieces.append(lines[i].strip().lstrip("> ").strip())
                i += 1
            yield "quote", " ".join(pieces)
            continue

        pieces = [stripped]
        i += 1
        while i < n:
            nxt = lines[i]
            if not nxt.strip():
                break
            if (
                HEADING_RE.match(nxt.strip())
                or nxt.strip().startswith("|")
                or nxt.strip().startswith("$$")
                or nxt.strip().startswith("![")
                or nxt.strip() in {"---", "***"}
                or LIST_RE.match(nxt)
            ):
                break
            pieces.append(nxt.strip())
            i += 1
        yield "para", " ".join(pieces)


def fill_cover(document: Document) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text
        if text.startswith("MÁSTER EN"):
            replace_paragraph_text(paragraph, COVER["master"])
        elif "TÍTULO DE TU TRABAJO" in text or "TITULO DE TU TRABAJO" in text:
            replace_paragraph_text(paragraph, TITLE)
        elif text.startswith("TFM elaborado por:"):
            replace_paragraph_text(paragraph, f"TFM elaborado por: {COVER['author']}")
        elif text.startswith("Tutor"):
            replace_paragraph_text(paragraph, f"Tutor/a de TFM: {COVER['tutor']}")
        elif "Ciudad y fecha" in text:
            replace_paragraph_text(paragraph, COVER["place_date"])

    # Subtítulo en el primer párrafo vacío posterior al título.
    after_title = False
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == TITLE:
            after_title = True
            continue
        if after_title and not paragraph.text.strip():
            if paragraph.runs:
                paragraph.runs[0].text = SUBTITLE
                for extra in paragraph.runs[1:]:
                    extra.text = ""
            else:
                paragraph.add_run(SUBTITLE)
            for extra_run in paragraph.runs:
                extra_run.italic = True
                set_run_font(extra_run, "Roboto Light", 16)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break


def add_preliminaries(document: Document) -> None:
    document.add_page_break()
    document.add_page_break()

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(18)
    heading.paragraph_format.space_before = Pt(72)
    run = heading.add_run("Agradecimientos")
    set_run_font(run, "Calibri", 18, True)

    body = document.add_paragraph()
    set_paragraph_body(body)
    run = body.add_run(
        "Apartado opcional según las normas de presentación. "
        "Redactar los agradecimientos en esta página o dejarla en blanco."
    )
    set_run_font(run, "Calibri", 11)
    run.italic = True


def _clear_paragraph_runs(paragraph: Paragraph) -> None:
    p_el = paragraph._p
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)


def write_page_footer(footer) -> None:
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    _clear_paragraph_runs(paragraph)
    add_page_number(paragraph)


def start_body_section(document: Document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    try:
        section.first_page_header.is_linked_to_previous = False
        section.first_page_footer.is_linked_to_previous = False
    except Exception:
        pass
    if section.header.paragraphs:
        section.header.paragraphs[0].text = ""
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.footer_distance = Cm(1.25)
    set_section_start(section, 4)
    write_page_footer(section.footer)
    try:
        write_page_footer(section.first_page_footer)
    except Exception:
        pass
    return section


def suppress_cover_page_numbers(document: Document) -> None:
    section = document.sections[0]
    section.footer.is_linked_to_previous = False
    if section.footer.paragraphs:
        section.footer.paragraphs[0].text = ""
    try:
        section.first_page_footer.is_linked_to_previous = False
        if section.first_page_footer.paragraphs:
            section.first_page_footer.paragraphs[0].text = ""
    except Exception:
        pass
    set_section_start(section, 1)


def set_core_properties(document: Document) -> None:
    now = datetime(2026, 9, 3, 22, 0, 0)
    cp = document.core_properties
    cp.author = "Raúl"
    cp.last_modified_by = "Raúl"
    cp.title = TITLE
    cp.subject = SUBTITLE
    cp.category = "Trabajo Fin de Máster"
    cp.keywords = "blanqueamiento coralino; aprendizaje automático; TFM"
    cp.comments = "Memoria de TFM. Septiembre de 2026."
    try:
        cp.created = now
    except Exception:
        pass
    cp.modified = datetime.now()
    try:
        cp.revision = 1
    except Exception:
        pass


def update_word_fields(path: Path) -> bool:
    """Actualiza el índice y los números de página del pie, sin reescribir el cuerpo."""
    escaped = str(path.resolve()).replace("'", "''")
    script = (
        "$ErrorActionPreference = 'Stop'\n"
        "$word = New-Object -ComObject Word.Application\n"
        "$word.Visible = $false\n"
        "$word.DisplayAlerts = 0\n"
        "try {\n"
        f"  $doc = $word.Documents.Open('{escaped}')\n"
        "  foreach ($toc in $doc.TablesOfContents) { $toc.Update() | Out-Null }\n"
        "  foreach ($section in $doc.Sections) {\n"
        "    foreach ($footer in $section.Footers) { $footer.Range.Fields.Update() | Out-Null }\n"
        "  }\n"
        "  $doc.Save()\n"
        "  $doc.Close($true)\n"
        "} finally {\n"
        "  $word.Quit() | Out-Null\n"
        "}\n"
    )
    import subprocess

    result = subprocess.run(
        ["powershell", "-NoProfile", "-Command", script],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if result.returncode != 0:
        err = (result.stderr or result.stdout or "").strip()
        print("No se pudieron actualizar los campos de Word:", err[:500])
        return False
    return True


def persist_core_properties(path: Path) -> None:
    """Vuelve a escribir autor y fechas: Word COM deja dc:creator vacío al guardar."""
    document = Document(str(path))
    set_core_properties(document)
    document.save(str(path))


def render_chapter(document: Document, markdown: str, hanging_refs: bool) -> int:
    figures = 0
    for kind, payload in iter_blocks(markdown):
        if kind == "heading":
            level, title = payload
            add_heading(document, title, level)
        elif kind == "image":
            add_image(document, payload)
            figures += 1
        elif kind == "caption":
            cleaned = strip_md_inline(payload)
            add_caption(document, cleaned)
        elif kind == "equation":
            add_equation(document, payload)
        elif kind == "table":
            add_table(document, payload)
        elif kind == "list":
            ordered, level, text = payload
            add_list_item(document, text, ordered, level)
        elif kind == "quote":
            paragraph = add_body_paragraph(document, payload)
            paragraph.paragraph_format.left_indent = Cm(1.0)
            for run in paragraph.runs:
                run.italic = True
        elif kind == "para":
            hanging = hanging_refs and bool(REF_ENTRY_RE.match(payload))
            add_body_paragraph(document, payload, hanging=hanging)
    return figures


def build() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"No se encuentra la plantilla de portada: {TEMPLATE}")

    document = Document(str(TEMPLATE))
    configure_styles(document)
    set_core_properties(document)
    fill_cover(document)
    suppress_cover_page_numbers(document)
    add_preliminaries(document)
    start_body_section(document)

    index_title = document.add_paragraph()
    index_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    index_title.paragraph_format.space_after = Pt(12)
    index_title.paragraph_format.space_before = Pt(0)
    run = index_title.add_run("Índice")
    set_run_font(run, "Calibri", 18, True)
    add_toc(document)

    total_figures = 0
    for name in CHAPTERS:
        path = REPORTS_DIR / name
        text = path.read_text(encoding="utf-8").strip()
        text, _ = embed_figures(text)
        total_figures += render_chapter(
            document,
            text,
            hanging_refs=name.startswith("08_"),
        )

    enable_update_fields(document)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT_PATH))
    if update_word_fields(OUTPUT_PATH):
        print("Campos de Word actualizados (índice y numeración de página).")
    persist_core_properties(OUTPUT_PATH)

    size_kb = OUTPUT_PATH.stat().st_size / 1024
    print(f"Documento generado: {OUTPUT_PATH}")
    print(f"Figuras incrustadas : {total_figures}")
    print(f"Tamaño              : {size_kb:.1f} KB")
    print("Portada: complete titulación, autor y tutor en el propio documento Word.")


if __name__ == "__main__":
    build()
